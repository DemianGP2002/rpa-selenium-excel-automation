import os
import time
import datetime
import openpyxl
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import Select, WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager


def formatear_peso(valor):
    if isinstance(valor, (datetime.datetime, datetime.date)):
        return f"{valor.day}.{valor.month}"
    return str(valor or "").strip()


def ejecutar_automatizacion():
    archivo_excel = "Reporte_Diario_Gallinas.xlsx"

    if not os.path.exists(archivo_excel):
        print(f"[ERROR] No existe el archivo en la ruta: {archivo_excel}")
        return

    print(f"[PASO 1] Abriendo Excel '{archivo_excel}'...")
    wb = openpyxl.load_workbook(archivo_excel)
    hoja = wb.active
    print(f"  -> Hoja '{hoja.title}' con {hoja.max_row - 1} registros.")

    opciones = webdriver.ChromeOptions()
    opciones.add_argument("--start-maximized")
    driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=opciones)
    wait = WebDriverWait(driver, 10)

    try:
        ruta_portal = os.path.abspath("portal_granja.html")
        driver.get(f"file://{ruta_portal}")
        print("[PASO 2] Cargando pantalla de inicio de sesión...")
        time.sleep(1.0)

        campo_usuario = wait.until(EC.presence_of_element_located((By.ID, "username")))
        campo_usuario.send_keys("admin")
        time.sleep(0.6)

        campo_pass = driver.find_element(By.ID, "password")
        campo_pass.send_keys("granja2026")
        time.sleep(0.6)

        driver.find_element(By.ID, "btn-login").click()

        wait.until(EC.visibility_of_element_located((By.ID, "seccion-panel")))
        print("[EXITO] Sesión autenticada correctamente.\n")
        time.sleep(1.2)

        ultimo_folio = ""
        for fila in range(2, hoja.max_row + 1):
            celda_id = hoja.cell(row=fila, column=1).value
            celda_zona = hoja.cell(row=fila, column=2).value
            celda_peso = hoja.cell(row=fila, column=3).value
            celda_estado = hoja.cell(row=fila, column=4).value

            if not celda_id or str(celda_id).strip().lower() in ("", "none"):
                print(f"  [OMITIDO] Fila {fila}: Sin ID_Ave. Se anota error en Excel.")
                hoja.cell(row=fila, column=5).value = "ERROR_ID_VACIO"
                time.sleep(0.5)
                continue

            id_ave = str(celda_id).strip()
            zona = str(celda_zona).strip() if celda_zona else "Corral Norte"
            peso = formatear_peso(celda_peso)
            estado = str(celda_estado).strip() if celda_estado else "Sano"

            print(f"--- Fila {fila}: Ave={id_ave} | Zona={zona} | Peso={peso}kg | Estado={estado} ---")

            campo_id = wait.until(EC.element_to_be_clickable((By.ID, "id_ave")))
            campo_id.clear()
            campo_id.send_keys(id_ave)
            time.sleep(0.3)

            Select(driver.find_element(By.ID, "zona_actual")).select_by_visible_text(zona)
            time.sleep(0.3)

            campo_peso = driver.find_element(By.ID, "peso_kg")
            campo_peso.clear()
            campo_peso.send_keys(peso)
            time.sleep(0.3)

            Select(driver.find_element(By.ID, "estado_salud")).select_by_visible_text(estado)
            time.sleep(0.3)

            chk_alerta = driver.find_element(By.ID, "alerta_sanitaria")
            es_riesgo = "observac" in estado.lower() or "escape" in estado.lower()
            if es_riesgo and not chk_alerta.is_selected():
                chk_alerta.click()
                print("  -> Casilla de Alerta Sanitaria marcada.")
            elif not es_riesgo and chk_alerta.is_selected():
                chk_alerta.click()
                print("  -> Casilla de Alerta Sanitaria desmarcada.")
            time.sleep(0.4)

            driver.find_element(By.ID, "btn-guardar").click()

            wait.until(lambda d: d.find_element(By.ID, "folio-comprobante").text.strip() not in ("", ultimo_folio))
            comprobante = driver.find_element(By.ID, "folio-comprobante").text.strip()
            ultimo_folio = comprobante

            hoja.cell(row=fila, column=5).value = comprobante
            print(f"  -> Comprobante asignado: {comprobante}")
            time.sleep(0.8)

        wb.save(archivo_excel)
        wb.close()
        print("\n" + "=" * 60)
        print(f"[FINALIZADO] Excel '{archivo_excel}' guardado exitosamente.")
        print("=" * 60)

    except Exception as error:
        print(f"[ERROR CRÍTICO] Falló el proceso: {error}")

    finally:
        print("\nCerrando navegador en 3 segundos...")
        time.sleep(3.0)
        driver.quit()


if __name__ == "__main__":
    ejecutar_automatizacion()